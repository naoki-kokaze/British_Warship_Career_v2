"""v2 の LLM 判定者（LLM-as-a-Judge）：院生と同じ指示書（v2版）で、v2 の RDF を艦ごとに3軸で判定させる（2026-10-14）

方針（2026-10-14 ユーザー承認。ログ 10/14）
    - 院生と同じ版の指示書で判定させる。指示書は試しの判定（〜10/16）のすり合わせで確定し、確定した版で設定を固定してから
      本評価の艦を判定させる。本評価の院生の判定を見てプロンプトを直さない（直すと一致率が甘くなる。査読 A2）
    - 試しの5隻での動作確認・調整は可（試しの5隻は評価の数字に入れない）
    - 軸2は、どの艦でも全項目を確かめさせる（院生は「気づいた場合のみ」の艦が多い。軸2の比較は「全項目」の艦で行う）
    - 再現性：モデル ID・プロンプトの版・指示書とオントロジーのハッシュを記録し、生の応答を保存する。記録シートの形の CSV は
      保存した応答から決定論的に作る。原文（Winfield）は著作権のため生の応答には保存しない（入力は元のファイルから再現できる）
    - 8/16 のパイロット版（llm_judge_api.py、v1 の RDF・指示書 7/16 版）は記録として残し、触らない

使い方（サンドボックスの外で。API の認証情報が必要）
    JINMON_DATASET=eval_v2 python3 llm_judge_v2.py show SM_9036            # API を呼ばずに、入力の大きさと RDF の抜き出しを確かめる
    JINMON_DATASET=eval_v2 python3 llm_judge_v2.py run 1 --ships trial     # 試しの5隻。保存済みの艦は呼ばない（途中から再開できる）
    JINMON_DATASET=eval_v2 python3 llm_judge_v2.py run 2 --ships main      # 本評価30隻（院生と同じ艦。指示書10/19版・プロンプト v2_2）
    JINMON_DATASET=eval_v2 python3 llm_judge_v2.py sheet 1 [--ships trial|main|<ID,...>]   # 記録シートの形の CSV を作る（API は呼ばない）

入力: ../docs/jinmoncom2026_annotation_guideline_v2_20261008.md、../outputs/ontology_v2/warship_career_ontology_v2.ttl、
      prompts/llm_judge_v2.md、<データセット>/rdf_v2/header_v2.ttl・events_v2.ttl、原文（STRATA の text 列）、eval_v2_selection.csv
出力: <データセット>/llm_judge_v2/raw/run<N>/<艦>.json、judge_run<N>[_<絞り込み>].csv
依存: anthropic（0.104.1 で確認）、rdflib、openpyxl
"""
import csv
import glob
import hashlib
import json
import os
import re
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import dataset_paths as P   # 開発用／評価用の切り替え（JINMON_DATASET）

GUIDELINE = os.path.join(HERE, "../docs/jinmoncom2026_annotation_guideline_v2_20261008.md")
ONTOLOGY = os.path.join(P.ONTOLOGY_DIR, "warship_career_ontology_v2.ttl")
PROMPT = os.path.join(HERE, "prompts/llm_judge_v2.md")
# 比較用（案B、llm_only_v2.py）の RDF を同じ判定者で判定するときは、読み込み先と保存先だけを環境変数で切り替える
# （プロンプト・指示書・メッセージの形は同じ。どちらの出力かは判定者に知らせない）
RDF_DIR = os.environ.get("JUDGE_RDF_DIR") or P.out("rdf_v2")
SELECTION = P.out("eval_v2_selection.csv")
OUT = os.environ.get("JUDGE_OUT") or P.out("llm_judge_v2")
# 判定者のモデル。JUDGE_MODEL で切り替える（10/20：gemini-4-argon を2つ目の判定者に追加。LLM の出力に甘いのが
# Claude の判定者だけか、LLM の判定者全般かを確かめるため）。gemini で始まるものは Gemini API（思考の量は既定）
MODEL = os.environ.get("JUDGE_MODEL") or "claude-opus-5-5"
EFFORT = "medium"           # パイロット（Option B）と同じ（Claude のみ）
MAX_TOKENS = 16000          # 打ち切られたら 1.75 倍で1回だけ再試行（パイロットと同じ）
PROMPT_VERSION = "llm_judge_v2_2"   # 10/19：指示書10/19版（規則の凍結 1f96b2d の後）。迷った点の用途、【設計】の印
ERROR_TYPES = ["タイプ誤用", "プロパティ誤用", "日付誤り", "金額誤り", "場所誤り", "寸法誤り", "順序誤り",
               "ハルシネーション", "欠落", "その他"]
SCHEMA = {
    "type": "object",
    "properties": {
        "axis1": {"type": "string", "enum": ["○", "×"]},
        "axis1_comment": {"type": "string"},
        "axis2_overall": {"type": "string", "enum": ["○", "△", "×"]},
        "axis3_overall": {"type": "string", "enum": ["○", "△", "×"]},
        "findings": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "target_id": {"type": "string"},
                    "axis2_judgment": {"type": "string", "enum": ["", "△", "×"]},
                    "axis2_error_type": {"type": "string", "enum": [""] + ERROR_TYPES},
                    "axis2_comment": {"type": "string"},
                    "axis3_judgment": {"type": "string", "enum": ["", "△", "×"]},
                    "axis3_error_type": {"type": "string", "enum": [""] + ERROR_TYPES},
                    "axis3_comment": {"type": "string"},
                },
                "required": ["target_id", "axis2_judgment", "axis2_error_type", "axis2_comment",
                             "axis3_judgment", "axis3_error_type", "axis3_comment"],
                "additionalProperties": False,
            },
        },
        "overall_comment": {"type": "string"},
        "unclear_points": {"type": "string"},
    },
    "required": ["axis1", "axis1_comment", "axis2_overall", "axis3_overall", "findings", "overall_comment",
                 "unclear_points"],
    "additionalProperties": False,
}
# 記録シート（指示書 §6.2）と同じ列
SHEET_FIELDS = ["行タイプ", "vessel_id", "vessel_name", "系統", "一覧", "対象ID", "軸1_構文", "軸2_スキーマ", "軸2_エラータイプ",
                "軸2_コメント", "軸3_意味", "軸3_エラータイプ", "軸3_コメント", "総合コメント", "所要時間（分）", "迷った点",
                "アノテーター", "作業日", "軸2の確認範囲"]


def sha12(text):
    return hashlib.sha256(text.encode()).hexdigest()[:12]


def read(path):
    return open(path, encoding="utf-8").read()


def system_text():
    # str.format は使わない（オントロジーの Turtle に波括弧が出うるため）
    return read(PROMPT).replace("{guideline}", read(GUIDELINE)).replace("{ontology}", read(ONTOLOGY))


def selection():
    return {r["id"]: r for r in csv.DictReader(open(SELECTION, encoding="utf-8-sig"))}


def ship_filter(arg):
    """--ships の値（艦のIDのカンマ区切り、または選定一覧の role：trial / main / reserve）を艦の集合にする。
    main は院生と同じ本評価の艦（eval_v2_main_order.csv の必須・追加の30隻）。選定一覧の main 42隻すべては main_all"""
    if arg == "main":
        rows = csv.DictReader(open(P.out("eval_v2_main_order.csv"), encoding="utf-8-sig"))
        return {r["id"] for r in rows if r["tier"] in ("必須", "追加")}
    if arg == "main_all":
        arg = "main"
    if arg in ("trial", "main", "reserve"):
        return {i for i, r in selection().items() if r["role"] == arg}
    return set(arg.split(","))


def load_graph():
    from rdflib import Graph
    g = Graph()
    for fn in ("header_v2.ttl", "events_v2.ttl"):
        g.parse(os.path.join(RDF_DIR, fn))
    return g


def ship_turtle(g, ship):
    """その艦のノード（艦の一覧と同じ選び方：make_ship_views_v2.py）と、そこからたどれる空白ノード、
    参照している共有のノード（場所・通貨・単位など）の型とラベルを Turtle にする"""
    from rdflib import BNode, Graph, URIRef
    from rdflib.namespace import RDF, RDFS
    sub = Graph()
    for prefix, ns in g.namespace_manager.namespaces():
        sub.bind(prefix, ns)

    def mine(s):
        name = str(s).rsplit("/", 1)[-1]
        return name == ship or name.startswith(ship + "_") or name.endswith("_" + ship) or ("_" + ship + "_") in name

    frontier = [s for s in set(g.subjects()) if isinstance(s, URIRef) and mine(s)]
    seen, shared = set(), set()
    while frontier:
        s = frontier.pop()
        if s in seen:
            continue
        seen.add(s)
        for p, o in g.predicate_objects(s):
            sub.add((s, p, o))
            if isinstance(o, BNode):
                frontier.append(o)
            elif isinstance(o, URIRef) and "/myData/" in str(o) and not mine(o):
                shared.add(o)
    for o in shared:
        for p in (RDF.type, RDFS.label):
            for v in g.objects(o, p):
                sub.add((o, p, v))
    return sub.serialize(format="turtle")


def source_text(ship):
    import openpyxl
    prefix = ship.rsplit("_", 1)[0]
    wb = openpyxl.load_workbook(os.path.join(P.STRATA, f"{prefix}.xlsx"), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    head = list(next(rows))
    for row in rows:
        if row[head.index("ID")] == f"myData:{ship}":
            return str(row[head.index("text")] or "")
    return None


def user_message(ship, sel, text, turtle):
    return (f"対象の艦: {ship}（{sel.get('name', '')}、系統 {sel.get('group', '')} {sel.get('group_name', '')}）\n\n"
            f"## 原文（Winfield 2014、OCR のテキスト）\n\n```\n{text}\n```\n\n"
            f"## RDF（header_v2.ttl・events_v2.ttl からこの艦の分を抜き出したもの）\n\n```turtle\n{turtle}\n```\n\n"
            "指示書とオントロジーに照らして判定し、指定の JSON スキーマで返してください。")


def show(ship):
    g = load_graph()
    sel = selection().get(ship, {})
    turtle, text, system = ship_turtle(g, ship), source_text(ship), system_text()
    msg = user_message(ship, sel, text or "", turtle)
    print(f"system: {len(system)} 文字（指示書 {sha12(read(GUIDELINE))}・オントロジー {sha12(read(ONTOLOGY))}）")
    print(f"user: {len(msg)} 文字（原文 {len(text or '')} 文字、RDF {len(turtle)} 文字・{turtle.count(' .')} 文の目安）")
    print(turtle)


def call_gemini(client, system, msg, max_tokens):
    """Gemini で1隻を判定する。Claude の記録と同じ項目を返す（stop_reason は STOP を end_turn に読み替える）"""
    import logging
    from google.genai import types
    logging.getLogger("google_genai.types").setLevel(logging.ERROR)   # 思考の署名の警告（本文の取り出しには影響しない）
    for attempt in range(4):
        try:
            resp = client.models.generate_content(
                model=MODEL, contents=msg,
                config=types.GenerateContentConfig(system_instruction=system, max_output_tokens=max_tokens,
                                                   response_mime_type="application/json", response_json_schema=SCHEMA))
            break
        except Exception as e:
            msg_e = str(e)
            print(f"  retry {attempt + 1}: {msg_e[:1500]}", file=sys.stderr)
            if "limit: 0" in msg_e or "NOT_FOUND" in msg_e:
                return None
            time.sleep(30 * (attempt + 1))
    else:
        return None
    cand = resp.candidates[0] if resp.candidates else None
    finish = str(cand.finish_reason) if cand else "no candidate"
    u = resp.usage_metadata
    return {"served_model": resp.model_version, "effort": "default", "max_tokens": max_tokens,
            "stop_reason": "end_turn" if finish.endswith("STOP") else finish, "finish_reason_raw": finish,
            "output_text": resp.text or "",
            "usage": {"input_tokens": u.prompt_token_count, "output_tokens": u.candidates_token_count,
                      "thoughts_tokens": u.thoughts_token_count, "cache_read_input_tokens": u.cached_content_token_count}}


def run(run_id, ships=None, limit=None):
    gemini = MODEL.startswith("gemini")
    if gemini:
        from google import genai
        client = genai.Client()
        limit_out = client.models.get(model=MODEL).output_token_limit
        gemini_max = min(32768, limit_out or 32768)
        print(f"judge {MODEL}: max_output_tokens {gemini_max}（モデルの上限 {limit_out}）")
    else:
        import anthropic
        client = anthropic.Anthropic()
    g = load_graph()
    sel = selection()
    system = system_text()
    meta = {"prompt_version": PROMPT_VERSION, "prompt_sha256_12": sha12(read(PROMPT)),
            "guideline_sha256_12": sha12(read(GUIDELINE)), "ontology_sha256_12": sha12(read(ONTOLOGY)),
            "system_sha256_12": sha12(system)}
    raw_dir = os.path.join(OUT, "raw", f"run{run_id}")
    os.makedirs(raw_dir, exist_ok=True)
    targets = sorted(ships if ships is not None else sel)
    done = 0
    for ship in targets:
        path = os.path.join(raw_dir, f"{ship}.json")
        if os.path.exists(path):
            continue
        if limit is not None and done >= limit:
            break
        text = source_text(ship)
        if text is None:
            print(f"{ship}: no source text", file=sys.stderr)
            continue
        turtle = ship_turtle(g, ship)
        msg = user_message(ship, sel.get(ship, {}), text, turtle)
        if gemini:
            got = call_gemini(client, system, msg, gemini_max)
            if got is None:
                print(f"{ship}: failed after retries", file=sys.stderr)
                continue
            record = {"ship": ship, "run": run_id, "requested_model": MODEL, **got, **meta,
                      "rdf_sha256_12": sha12(turtle), "source_sha256_12": sha12(text),
                      "time": time.strftime("%Y-%m-%dT%H:%M:%S%z")}
            json.dump(record, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            done += 1
            print(f"{ship}: {record['stop_reason']} ({got['usage']['input_tokens']}+{got['usage']['output_tokens']} tokens)")
            continue
        max_tokens, resp = MAX_TOKENS, None
        for attempt in range(4):
            try:
                resp = client.messages.create(
                    model=MODEL,
                    max_tokens=max_tokens,
                    system=[{"type": "text", "text": system, "cache_control": {"type": "ephemeral", "ttl": "1h"}}],
                    output_config={"effort": EFFORT, "format": {"type": "json_schema", "schema": SCHEMA}},
                    messages=[{"role": "user", "content": msg}],
                )
            except anthropic.RateLimitError:
                time.sleep(30 * (attempt + 1))
                continue
            except anthropic.APIConnectionError:
                time.sleep(10 * (attempt + 1))
                continue
            if resp.stop_reason == "max_tokens" and max_tokens == MAX_TOKENS:
                max_tokens = int(MAX_TOKENS * 1.75)
                continue
            break
        if resp is None:
            print(f"{ship}: failed after retries", file=sys.stderr)
            continue
        out_text = next((b.text for b in resp.content if b.type == "text"), "")
        record = {"ship": ship, "run": run_id, "requested_model": MODEL, "served_model": resp.model, "effort": EFFORT,
                  "max_tokens": max_tokens, "stop_reason": resp.stop_reason, **meta,
                  "rdf_sha256_12": sha12(turtle), "source_sha256_12": sha12(text), "output_text": out_text,
                  "usage": {"input_tokens": resp.usage.input_tokens, "output_tokens": resp.usage.output_tokens,
                            "cache_read_input_tokens": getattr(resp.usage, "cache_read_input_tokens", None),
                            "cache_creation_input_tokens": getattr(resp.usage, "cache_creation_input_tokens", None)},
                  "time": time.strftime("%Y-%m-%dT%H:%M:%S%z")}
        json.dump(record, open(path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        done += 1
        print(f"{ship}: {resp.stop_reason} ({resp.usage.input_tokens}+{resp.usage.output_tokens} tokens, "
              f"cache read {record['usage']['cache_read_input_tokens']})")
    print(f"run{run_id}: {done} ships judged")


WORST = {"": 0, "○": 0, "△": 1, "×": 2}


def sheet(run_id, ships=None, tag=""):
    sel = selection()
    rows, problems = [], []
    for path in sorted(glob.glob(os.path.join(OUT, "raw", f"run{run_id}", "*.json"))):
        rec = json.load(open(path, encoding="utf-8"))
        ship = rec["ship"]
        if ships is not None and ship not in ships:
            continue
        if rec["stop_reason"] != "end_turn":
            problems.append(f"{ship}: stop_reason {rec['stop_reason']}")
            continue
        try:
            j = json.loads(rec["output_text"])
        except json.JSONDecodeError:
            problems.append(f"{ship}: invalid JSON")
            continue
        s = sel.get(ship, {})
        for ax in ("axis2", "axis3"):   # 総合が明細の最も重い判定より軽いものを知らせる（直さない。記録のまま）
            worst = max((WORST[f[f"{ax}_judgment"]] for f in j["findings"]), default=0)
            if WORST[j[f"{ax}_overall"]] < worst:
                problems.append(f"{ship}: {ax}_overall {j[f'{ax}_overall']} is lighter than its findings")
        annotator = f"LLM:{rec['served_model']}:{rec['prompt_version']}"
        rows.append({"行タイプ": "総合判定", "vessel_id": ship, "vessel_name": s.get("name", ""),
                     "系統": f"{s.get('group', '')} {s.get('group_name', '')}".strip(), "一覧": f"{ship} {s.get('name', '')}",
                     "軸1_構文": j["axis1"], "軸2_スキーマ": j["axis2_overall"], "軸3_意味": j["axis3_overall"],
                     "総合コメント": "\n".join(x for x in (j["overall_comment"], j["axis1_comment"]) if x),
                     "迷った点": j["unclear_points"], "アノテーター": annotator, "作業日": rec["time"][:10],
                     "軸2の確認範囲": "全項目"})
        for f in j["findings"]:
            if not f["axis2_judgment"] and not f["axis3_judgment"]:
                continue
            rows.append({"行タイプ": "明細", "vessel_id": ship, "対象ID": f["target_id"],
                         "軸2_スキーマ": f["axis2_judgment"], "軸2_エラータイプ": f["axis2_error_type"],
                         "軸2_コメント": f["axis2_comment"], "軸3_意味": f["axis3_judgment"],
                         "軸3_エラータイプ": f["axis3_error_type"], "軸3_コメント": f["axis3_comment"],
                         "アノテーター": annotator, "作業日": rec["time"][:10]})
    path = os.path.join(OUT, f"judge_run{run_id}{'_' + tag if tag else ''}.csv")
    with open(path, "w", encoding="utf-8-sig", newline="") as fp:
        w = csv.DictWriter(fp, fieldnames=SHEET_FIELDS)
        w.writeheader()
        w.writerows({k: r.get(k, "") for k in SHEET_FIELDS} for r in rows)
    n = sum(r["行タイプ"] == "総合判定" for r in rows)
    print(f"{n} ships, {len(rows) - n} findings -> {path}")
    for p in problems:
        print("  " + p)


if __name__ == "__main__":
    args = sys.argv[1:]
    ships_arg = args[args.index("--ships") + 1] if "--ships" in args else None
    ships = ship_filter(ships_arg) if ships_arg else None
    if args[:1] == ["show"] and len(args) > 1:
        show(args[1])
    elif args[:1] == ["run"] and len(args) > 1:
        lim = int(args[args.index("--limit") + 1]) if "--limit" in args else None
        run(int(args[1]), ships, lim)
    elif args[:1] == ["sheet"] and len(args) > 1:
        sheet(int(args[1]), ships, re.sub(r"[^A-Za-z0-9_]", "", ships_arg or ""))
    else:
        print(__doc__)
